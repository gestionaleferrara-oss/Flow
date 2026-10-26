"""Download di vendite e acquisti per periodo (pagina Importa Dati)."""
import csv
import io
from datetime import datetime
from typing import Optional

from fastapi import APIRouter, HTTPException, Query
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from database import db

router = APIRouter()

# (campo nel database, intestazione nel file, tipo)
COLONNE_VENDITE = [
    ("data", "Data", "data"),
    ("cliente_nome", "Cliente", "testo"),
    ("prodotto_nome", "Prodotto", "testo"),
    ("quantita", "Quantità", "numero"),
    ("prezzo_unitario", "Prezzo unitario", "euro"),
    ("totale", "Totale", "euro"),
    ("costo_unitario_fifo", "Costo unitario (FIFO)", "euro"),
    ("costo_fifo", "Costo totale (FIFO)", "euro"),
    ("movimento", "Movimento", "testo"),
    ("causale", "Causale", "testo"),
]

COLONNE_ACQUISTI = [
    ("data", "Data", "data"),
    ("fornitore_nome", "Fornitore", "testo"),
    ("prodotto_nome", "Prodotto", "testo"),
    ("quantita", "Quantità", "numero"),
    ("kg_per_pezzo", "Kg per pezzo", "numero"),
    ("totale_kg", "Totale kg", "numero"),
    ("prezzo_unitario", "Prezzo unitario", "euro"),
    ("totale", "Totale", "euro"),
]

TIPI = {
    "vendite": ("vendite", "Vendite", COLONNE_VENDITE),
    "acquisti": ("ordini", "Acquisti", COLONNE_ACQUISTI),
}


def _valida_data(valore: Optional[str], nome: str) -> Optional[str]:
    if not valore:
        return None
    try:
        datetime.strptime(valore, "%Y-%m-%d")
    except ValueError:
        raise HTTPException(status_code=400, detail=f"{nome} non valida: usa il formato AAAA-MM-GG")
    return valore


async def _leggi(collezione: str, colonne, start: Optional[str], end: Optional[str]):
    query = {}
    if start or end:
        query["data"] = {}
        if start:
            query["data"]["$gte"] = start
        if end:
            query["data"]["$lte"] = end
    proiezione = {"_id": 0, **{campo: 1 for campo, _, _ in colonne}}
    cursore = db[collezione].find(query, proiezione).sort("data", 1)
    return [doc async for doc in cursore]


def _valore(doc, campo, tipo):
    v = doc.get(campo)
    if v is None or v == "":
        return None
    if tipo == "data":
        try:
            return datetime.strptime(str(v)[:10], "%Y-%m-%d").date()
        except ValueError:
            return v
    if tipo in ("numero", "euro"):
        try:
            return float(v)
        except (TypeError, ValueError):
            return v
    return v


def _foglio(ws, titolo, colonne, righe):
    ws.title = titolo
    ws.append([intestazione for _, intestazione, _ in colonne])
    for cella in ws[1]:
        cella.font = Font(bold=True, color="FFFFFF")
        cella.fill = PatternFill("solid", fgColor="1F2937")
    for doc in righe:
        ws.append([_valore(doc, campo, tipo) for campo, _, tipo in colonne])
    for i, (_, intestazione, tipo) in enumerate(colonne, start=1):
        lettera = get_column_letter(i)
        formato = {"data": "DD/MM/YYYY", "euro": "#,##0.00", "numero": "#,##0.###"}.get(tipo)
        if formato:
            for (cella,) in ws.iter_rows(min_row=2, min_col=i, max_col=i):
                cella.number_format = formato
        larghezza = 40 if intestazione in ("Cliente", "Fornitore", "Prodotto") else 14
        ws.column_dimensions[lettera].width = larghezza
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def _csv(colonne, righe) -> bytes:
    out = io.StringIO()
    writer = csv.writer(out, delimiter=";")
    writer.writerow([intestazione for _, intestazione, _ in colonne])
    for doc in righe:
        riga = []
        for campo, _, tipo in colonne:
            v = _valore(doc, campo, tipo)
            if v is None:
                riga.append("")
            elif tipo == "data" and hasattr(v, "strftime"):
                riga.append(v.strftime("%d/%m/%Y"))
            elif isinstance(v, float):
                riga.append(f"{v:.4f}".rstrip("0").rstrip(".").replace(".", ","))
            else:
                riga.append(v)
        writer.writerow(riga)
    # BOM per far leggere correttamente gli accenti a Excel
    return ("﻿" + out.getvalue()).encode("utf-8")


@router.get("/export/movimenti")
async def export_movimenti(
    tipo: str = Query("entrambi", pattern="^(vendite|acquisti|entrambi)$"),
    formato: str = Query("xlsx", pattern="^(xlsx|csv)$"),
    start_date: Optional[str] = None,
    end_date: Optional[str] = None,
):
    start = _valida_data(start_date, "Data di inizio")
    end = _valida_data(end_date, "Data di fine")
    if start and end and start > end:
        raise HTTPException(status_code=400, detail="La data di inizio è successiva alla data di fine")

    tipi = ["vendite", "acquisti"] if tipo == "entrambi" else [tipo]
    if formato == "csv" and len(tipi) > 1:
        raise HTTPException(status_code=400, detail="Il CSV contiene un solo tipo: scegli Vendite o Acquisti, oppure usa Excel")

    periodo = f"{start or 'inizio'}_{end or 'oggi'}" if (start or end) else "tutto"
    nome_file = f"flow_{tipo}_{periodo}.{formato}"

    if formato == "csv":
        collezione, _, colonne = TIPI[tipi[0]]
        righe = await _leggi(collezione, colonne, start, end)
        contenuto = _csv(colonne, righe)
        media = "text/csv; charset=utf-8"
    else:
        wb = Workbook()
        for i, t in enumerate(tipi):
            collezione, titolo, colonne = TIPI[t]
            righe = await _leggi(collezione, colonne, start, end)
            ws = wb.active if i == 0 else wb.create_sheet()
            _foglio(ws, titolo, colonne, righe)
        buffer = io.BytesIO()
        wb.save(buffer)
        contenuto = buffer.getvalue()
        media = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

    return StreamingResponse(
        io.BytesIO(contenuto),
        media_type=media,
        headers={"Content-Disposition": f'attachment; filename="{nome_file}"'},
    )
